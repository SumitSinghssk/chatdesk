import io
import json

from openpyxl import Workbook, load_workbook

from tests.test_contacts import BASE, WEB, biz  # noqa: F401 - the fixture

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx(rows: list[list]) -> bytes:
    book = Workbook()
    ws = book.active
    for row in rows:
        ws.append(row)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _preview(biz, content: bytes, name="list.xlsx", **form):  # noqa: F811
    res = biz.post(f"{BASE}/import/preview", headers=WEB, data=form,
                   files={"file": (name, content, XLSX if name.endswith("xlsx") else "text/csv")})
    assert res.status_code == 200, res.text
    return res.json()["data"]


def test_excel_preview_marks_every_row(biz):  # noqa: F811
    biz.post(BASE, headers=WEB, json={"phone": "9876500001", "name": "Old Friend"})
    biz.post(BASE, headers=WEB, json={"phone": "9876500002", "name": "Left Us", "opted_out": True})
    content = _xlsx([
        [None],                                                     # a blank line above the header is fine
        ["Mobile No.", "Customer Name", "E-mail", "City", "Order ID"],
        [9876500001, "Old Friend", "friend@example.com", "Delhi", "A1"],   # number cell, already a contact
        [9876500003.0, "New Person", "not-an-email", "Pune", "A2"],        # float cell, bad email
        ["+91 98765 00003", "Same Again", "", "", ""],                     # duplicate of the row above
        ["9.17E+11", "Shortened", "", "", ""],                             # Excel scientific notation
        ["", "No Number", "", "", ""],
        ["12345", "Too Short", "", "", ""],
        ["9876500002", "Left Us", "", "", ""],                             # opted out
    ])
    data = _preview(biz, content)
    assert data["mapping"] == {"phone": "Mobile No.", "name": "Customer Name", "email": "E-mail", "tags": None,
                               "skip": []}
    assert [c["column"] for c in data["custom"]] == ["City", "Order ID"]
    assert data["custom"][1]["field"] == "order_id"
    by_row = {r["row"]: r for r in data["rows"]}
    assert by_row[3]["status"] == "update" and by_row[3]["wa_id"] == "919876500001"
    assert by_row[4]["status"] == "new" and by_row[4]["email"] is None and "isn't valid" in by_row[4]["warnings"][0]
    assert by_row[4]["fields"] == {"city": "Pune", "order_id": "A2"}
    assert by_row[5]["status"] == "duplicate" and by_row[5]["duplicate_of"] == 4
    assert by_row[6]["status"] == "invalid" and "shortened" in by_row[6]["message"]
    assert by_row[7]["message"] == "No phone number" and by_row[8]["status"] == "invalid"
    assert by_row[9]["status"] == "update" and by_row[9]["warnings"] == ["Opted out: won't get campaigns"]
    assert data["counts"] == {"total": 7, "new": 1, "update": 2, "exists": 0, "duplicate": 1, "invalid": 3,
                              "warnings": 2}

    # nothing was saved by the preview
    assert biz.get(BASE, params={"search": "New Person", "page": 1}).json()["data"] == []

    # turning "update existing" off leaves known numbers alone
    kept = _preview(biz, content, update_existing="false")
    assert kept["counts"]["exists"] == 2 and kept["counts"]["update"] == 0


def test_import_with_chosen_columns_and_tag(biz):  # noqa: F811
    content = _xlsx([
        ["Ref", "Contact", "Who", "Note"],
        ["R1", "+91 98111 22233", "Kavya Rao", "call after 6"],
        ["R2", "98111 22244", "Imran Ali", ""],
    ])
    # nothing looks like a "phone" header, but the numbers give the column away
    guessed = _preview(biz, content)
    assert guessed["mapping"]["phone"] == "Contact" and guessed["mapping"]["name"] is None

    mapping = json.dumps({"phone": "Contact", "name": "Who", "skip": ["Note"]})
    res = biz.post(f"{BASE}/import", headers=WEB, data={"mapping": mapping, "tag": "Sheet: Oct list",
                                                         "update_existing": "true"},
                   files={"file": ("oct.xlsx", content, XLSX)})
    assert res.status_code == 200, res.text
    data = res.json()["data"]
    assert (data["created"], data["updated"], data["skipped"]) == (2, 0, 0)
    assert data["tag"] == "Sheet: Oct list" and data["columns"]["custom"] == ["Ref"]
    kavya = biz.get(BASE, params={"search": "Kavya", "page": 1}).json()["data"][0]
    assert kavya["wa_id"] == "919811122233" and kavya["tags"] == ["Sheet: Oct list"]
    assert kavya["attributes"] == {"ref": "R1"}  # Note was skipped

    # no phone column at all: the preview still shows the sheet so the column can be picked
    unknown = _preview(biz, content, mapping=json.dumps({"phone": None}))
    assert unknown["mapping"]["phone"] is None and unknown["counts"]["invalid"] == 2
    assert unknown["rows"][0]["message"] == "Choose the column that has the phone numbers"
    refused = biz.post(f"{BASE}/import", headers=WEB, data={"mapping": json.dumps({"phone": None})},
                       files={"file": ("oct.xlsx", content, XLSX)})
    assert refused.status_code == 422 and "mapping.phone" in refused.json()["errors"]

    bad = biz.post(f"{BASE}/import/preview", headers=WEB, data={"mapping": json.dumps({"phone": "Nope"})},
                   files={"file": ("oct.xlsx", content, XLSX)})
    assert bad.status_code == 422 and "mapping.phone" in bad.json()["errors"]
    twice = biz.post(f"{BASE}/import/preview", headers=WEB,
                     data={"mapping": json.dumps({"phone": "Contact", "name": "Contact"})},
                     files={"file": ("oct.xlsx", content, XLSX)})
    assert twice.status_code == 422


def test_files_that_cannot_be_used(biz):  # noqa: F811
    old = biz.post(f"{BASE}/import/preview", headers=WEB, files={"file": ("old.xls", b"\xd0\xcf\x11\xe0", "x")})
    assert old.status_code == 422 and ".xlsx" in old.json()["errors"]["file"][0]
    empty = biz.post(f"{BASE}/import/preview", headers=WEB, files={"file": ("e.csv", b"\n\n", "text/csv")})
    assert empty.status_code == 422
    broken = biz.post(f"{BASE}/import/preview", headers=WEB, files={"file": ("b.xlsx", b"PK not a zip", XLSX)})
    assert broken.status_code == 422
    pdf = biz.post(f"{BASE}/import/preview", headers=WEB, files={"file": ("a.pdf", b"%PDF-1.4", "application/pdf")})
    assert pdf.status_code == 422


def test_semicolon_csv_and_repeated_headers(biz):  # noqa: F811
    content = "phone;name;name\n9811100001;A;B\n".encode("cp1252")
    data = _preview(biz, content, name="list.csv")
    assert data["headers"] == ["phone", "name", "name (2)"]
    assert data["rows"][0]["name"] == "A" and data["rows"][0]["fields"] == {"name_(2)": "B"}


def test_sample_files(biz):  # noqa: F811
    xlsx = biz.get(f"{BASE}/import/sample", params={"format": "xlsx"})
    assert xlsx.status_code == 200 and xlsx.headers["content-type"].startswith(XLSX)
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Contacts", "How to fill"] and book["Contacts"]["A1"].value == "Phone"
    csv_file = biz.get(f"{BASE}/import/sample", params={"format": "csv"})
    assert csv_file.status_code == 200 and csv_file.text.lstrip("﻿").startswith("Phone,Name")

    # the sample imports cleanly
    data = _preview(biz, xlsx.content, name="contacts-sample.xlsx")
    assert data["counts"]["new"] == 3 and data["counts"]["invalid"] == 0
