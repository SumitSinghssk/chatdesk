# ruff: noqa: F811  (the `business` fixture is imported from test_inbox and used as a parameter)
import http.server
import io
import json
import threading

import pytest

from app.core.database import SessionLocal
from app.modules.customer.ai_agent import engine as ai_engine
from app.modules.customer.ai_agent.models import AiThread
from app.modules.customer.automation import engine as auto_engine
from app.modules.customer.chatbots import engine as bot_engine
from tests import ai_fake
from tests.test_inbox import BASE, WEB, FakeMeta, business, first_conversation, inbound  # noqa: F401

AI = f"{BASE}/ai-agent"
PUB = "/api/v1/website/ai"


@pytest.fixture(scope="module")
def fake_ai():
    server = ai_fake.start()
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


@pytest.fixture()
def shop(business, fake_ai, monkeypatch):
    monkeypatch.setenv("AI_PROVIDER_BASE_URL", fake_ai)
    monkeypatch.setattr(auto_engine, "WhatsAppClient", FakeMeta)
    monkeypatch.setattr(bot_engine, "WhatsAppClient", FakeMeta)
    monkeypatch.setattr(ai_engine, "WhatsAppClient", FakeMeta)
    ai_fake.Fake.requests, ai_fake.Fake.fail, ai_fake.Fake.reply = [], None, None
    return business


def sse(res) -> list[tuple[str, dict]]:
    assert res.status_code == 200, res.text
    out = []
    for block in res.text.strip().split("\n\n"):
        lines = dict(line.split(": ", 1) for line in block.split("\n") if ": " in line)
        out.append((lines["event"], json.loads(lines["data"])))
    return out


def ask(client, agent_id, text, thread_id=None, channel="web"):
    events = sse(client.post(f"{AI}/agents/{agent_id}/test", headers=WEB,
                             json={"message": text, "thread_id": thread_id, "channel": channel}))
    thread = events[0][1]["id"]
    deltas = "".join(d["text"] for k, d in events if k == "delta")
    end = events[-1]
    return thread, deltas, end


def connect(client, provider="anthropic", key="sk-ant-test-key-1234567890", **extra):
    res = client.post(f"{AI}/connections", headers=WEB, json={"provider": provider, "api_key": key, **extra})
    assert res.status_code == 200, res.text
    return res.json()["data"]


def agent_with_knowledge(client, **behaviour):
    connect(client)
    agent = client.post(f"{AI}/agents", headers=WEB, json={"name": "Asha", "business_name": "Sumit Sweets"}).json()
    agent = agent["data"]
    a = agent["id"]
    assert client.post(f"{AI}/agents/{a}/sources/faq", headers=WEB, json={
        "question": "What are your timings?", "answer": "We are open from 9 am to 9 pm, all seven days."}).status_code == 200
    assert client.post(f"{AI}/agents/{a}/sources/text", headers=WEB, json={
        "title": "Delivery", "content": "Delivery is free above Rs 500 within Dehradun city."}).status_code == 200
    if behaviour:
        full = client.get(f"{AI}/agents/{a}").json()["data"]
        body = {"name": full["name"], "connection_id": full["connection"]["id"], "model": full["model"],
                "behaviour": {**full["behaviour"], **behaviour}, "limits": full["limits"]}
        assert client.put(f"{AI}/agents/{a}", headers=WEB, json=body).status_code == 200
    trained = client.post(f"{AI}/agents/{a}/train", headers=WEB)
    assert trained.status_code == 200, trained.text
    return a


def settings(client, a, **limits):
    full = client.get(f"{AI}/agents/{a}").json()["data"]
    body = {"name": full["name"], "connection_id": full["connection"]["id"], "model": full["model"],
            "behaviour": full["behaviour"], "limits": {**full["limits"], **limits}}
    res = client.put(f"{AI}/agents/{a}", headers=WEB, json=body)
    assert res.status_code == 200, res.text


def channels(client, a, **body):
    full = client.get(f"{AI}/agents/{a}").json()["data"]
    payload = {"whatsapp_enabled": full["whatsapp_enabled"], "whatsapp": full["whatsapp"],
               "widget_enabled": full["widget_enabled"], "widget": full["widget"], "api_enabled": full["api_enabled"]}
    for key, value in body.items():
        payload[key] = {**payload[key], **value} if isinstance(value, dict) else value
    res = client.put(f"{AI}/agents/{a}/channels", headers=WEB, json=payload)
    assert res.status_code == 200, res.text
    return res.json()["data"]


def texts_to(wa_id):
    return [s["text"] if isinstance(s["text"], str) else s["text"]["body"]
            for s in FakeMeta.sent if s.get("to") == wa_id and s.get("type") == "text"]


def pdf_bytes(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
               b"/Resources << /Font << /F1 5 0 R >> >> >>",
               b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = b"%PDF-1.4\n", []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


# Connections -------------------------------------------------------------------------------------

def test_connections_check_keys_and_never_show_them(shop):
    bad = shop.post(f"{AI}/connections", headers=WEB, json={"provider": "anthropic", "api_key": "bad"}).json()
    assert bad["data"]["status"] == "failed" and "rejected" in bad["data"]["error"]
    good = connect(shop)
    assert good["status"] == "ok" and good["models"][:3] == ["claude-opus-5-5", "claude-sonnet-5-5", "claude-haiku-5-5"]
    assert "test-key" not in json.dumps(good) and good["key_hint"].startswith("sk-ant-")
    openai = connect(shop, "openai", "sk-openai-0000000000000")
    assert openai["models"] == ["gpt-test", "gpt-test-mini"]  # embedding models left out
    custom = shop.post(f"{AI}/connections", headers=WEB, json={"provider": "custom", "api_key": "x"})
    assert custom.status_code == 422 and "base_url" in custom.json()["errors"]
    # editing without a new key keeps the old one
    edited = shop.put(f"{AI}/connections/{good['id']}", headers=WEB, json={"provider": "anthropic", "label": "Main"})
    assert edited.json()["data"]["label"] == "Main" and edited.json()["data"]["status"] == "ok"
    providers = shop.get(f"{AI}/providers").json()["data"]
    assert {p["key"] for p in providers} >= {"anthropic", "openai", "gemini", "groq", "openrouter", "custom"}
    listed = shop.get(f"{AI}/connections").json()["data"]
    assert len(listed) == 3
    assert shop.delete(f"{AI}/connections/{bad['data']['id']}", headers=WEB).status_code == 200


# Knowledge ----------------------------------------------------------------------------------------

def test_files_are_read_once_trained_and_previewed(shop):
    import docx
    import openpyxl

    connect(shop)
    a = shop.post(f"{AI}/agents", headers=WEB, json={"name": "Docs"}).json()["data"]["id"]
    doc = docx.Document()
    doc.add_heading("Return policy", 1)
    doc.add_paragraph("Sweets can be returned within 2 days if the box is unopened.")
    buf = io.BytesIO()
    doc.save(buf)
    book = openpyxl.Workbook()
    book.active.append(["Item", "Price"])
    book.active.append(["Kaju katli", "Rs 900 per kg"])
    xbuf = io.BytesIO()
    book.save(xbuf)
    files = {
        "menu.pdf": pdf_bytes("Gulab jamun costs Rs 400 per kg"),
        "policy.docx": buf.getvalue(),
        "prices.xlsx": xbuf.getvalue(),
        "stores.csv": b"City,Address\nDehradun,Rajpur Road\n",
        "notes.md": "# Parking\nFree parking behind the shop.".encode(),
    }
    ids = {}
    for name, content in files.items():
        res = shop.post(f"{AI}/agents/{a}/sources/file", headers=WEB, files={"file": (name, content)})
        assert res.status_code == 200, (name, res.text)
        ids[name] = res.json()["data"]["id"]
    sources = {s["filename"]: s for s in shop.get(f"{AI}/agents/{a}/sources").json()["data"]}
    assert "Gulab jamun" in sources["menu.pdf"]["excerpt"] and sources["menu.pdf"]["pages"] == 1
    assert "Return policy" in sources["policy.docx"]["excerpt"]
    assert "Item: Kaju katli; Price: Rs 900 per kg" in sources["prices.xlsx"]["excerpt"]
    assert "City: Dehradun; Address: Rajpur Road" in sources["stores.csv"]["excerpt"]
    assert shop.get(f"{AI}/agents/{a}").json()["data"]["knowledge"]["needs_training"] is True

    bad = shop.post(f"{AI}/agents/{a}/sources/file", headers=WEB, files={"file": ("x.exe", b"MZ")})
    assert bad.status_code == 422 and "PDF" in bad.json()["message"]
    fake = shop.post(f"{AI}/agents/{a}/sources/file", headers=WEB, files={"file": ("x.pdf", b"not a pdf")})
    assert fake.status_code == 422

    trained = shop.post(f"{AI}/agents/{a}/train", headers=WEB).json()
    assert trained["data"]["sources"] == 5 and trained["data"]["agent"]["knowledge"]["version"] == 1
    assert trained["data"]["agent"]["knowledge"]["mode"] == "full" and "Trained on 5 sources" in trained["message"]
    preview = shop.get(f"{AI}/agents/{a}/sources/{ids['notes.md']}").json()["data"]
    assert preview["content"].startswith("# Parking") and preview["chunks"][0]["heading"] == "Parking"
    original = shop.get(f"{AI}/agents/{a}/sources/{ids['menu.pdf']}/file")
    assert original.status_code == 200 and original.headers["content-type"] == "application/pdf"
    assert original.content.startswith(b"%PDF") and "inline" in original.headers["content-disposition"]
    hits = shop.get(f"{AI}/agents/{a}/search", params={"q": "kaju price"}).json()["data"]
    assert hits and hits[0]["title"] == "prices"

    # switching a source off leaves it out of the next training
    shop.put(f"{AI}/agents/{a}/sources/{ids['stores.csv']}", headers=WEB, json={"enabled": False})
    again = shop.post(f"{AI}/agents/{a}/train", headers=WEB).json()["data"]
    assert again["sources"] == 4 and again["agent"]["knowledge"]["version"] == 2
    assert shop.delete(f"{AI}/agents/{a}/sources/{ids['menu.pdf']}", headers=WEB).status_code == 200
    assert shop.get(f"{AI}/agents/{a}/sources/{ids['menu.pdf']}/file").status_code == 404


def test_website_pages_and_crawl(shop, monkeypatch):
    monkeypatch.setenv("CHATBOT_ALLOW_PRIVATE_URLS", "1")
    pages = {
        "/": "<html><head><title>Sumit Sweets</title></head><body><nav>Menu</nav><h1>Welcome</h1>"
             "<p>Fresh sweets made every morning in Dehradun.</p><a href='/about'>About</a>"
             "<a href='/login'>Login</a><a href='https://other.example.com/'>x</a></body></html>",
        "/about": "<html><head><title>About us</title></head><body><p>Family business since 1975, "
                  "now run by the third generation.</p></body></html>",
    }

    class Site(http.server.BaseHTTPRequestHandler):
        def log_message(self, *a):
            pass

        def do_GET(self):
            body = pages.get(self.path)
            self.send_response(200 if body else 404)
            self.send_header("Content-Type", "text/html")
            self.end_headers()
            self.wfile.write((body or "nope").encode())

    server = http.server.ThreadingHTTPServer(("127.0.0.1", 0), Site)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        connect(shop)
        a = shop.post(f"{AI}/agents", headers=WEB, json={"name": "Site"}).json()["data"]["id"]
        url = f"http://127.0.0.1:{server.server_address[1]}/"
        res = shop.post(f"{AI}/agents/{a}/sources/url", headers=WEB, json={"url": url, "crawl": True})
        assert res.status_code == 200, res.text
        assert res.json()["data"]["title"] == "Sumit Sweets" and "1 more" in res.json()["message"]
        assert "Menu" not in res.json()["data"]["excerpt"]  # navigation left out
        rows = shop.get(f"{AI}/agents/{a}/sources").json()["data"]
        assert {r["title"] for r in rows} == {"Sumit Sweets", "About us"}
        assert all(r["status"] == "ready" for r in rows)
        missing = shop.post(f"{AI}/agents/{a}/sources/url", headers=WEB, json={"url": url + "nothing"})
        assert missing.status_code == 422 and "404" in missing.json()["message"]
    finally:
        server.shutdown()
    monkeypatch.delenv("CHATBOT_ALLOW_PRIVATE_URLS")
    from app.core.config import settings

    monkeypatch.setattr(type(settings), "is_production", property(lambda self: True))
    private = shop.post(f"{AI}/agents/{a}/sources/url", headers=WEB, json={"url": "http://127.0.0.1:9/"})
    assert private.status_code == 422 and "public" in private.json()["message"]


# Answering -------------------------------------------------------------------------------------------

def test_test_chat_streams_answers_from_knowledge_with_caching(shop):
    a = agent_with_knowledge(shop)
    thread, streamed, (kind, done) = ask(shop, a, "What are your timings?")
    assert kind == "done" and done["outcome"] == "answered"
    assert done["text"] == "We are open from 9 am to 9 pm, all seven days." == streamed
    assert done["sources"] and done["sources"][0]["title"] == "What are your timings?"
    assert done["mode"] == "full" and done["model"] == "claude-opus-5-5"
    req = ai_fake.Fake.requests[-1]["body"]
    assert req["system"][0]["cache_control"] == {"type": "ephemeral"}  # the fixed part is cached
    assert "9 am to 9 pm" in req["system"][0]["text"] and "Now:" in req["system"][1]["text"]
    assert req["output_config"] == {"effort": "low"} and req["fallbacks"] == "default"

    # the same fixed prompt again: served from the provider's cache
    _, _, (_, second) = ask(shop, a, "Is delivery free?", thread)
    assert "free above Rs 500" in second["text"] and second["tokens"]["cached"] > 0 and second["saved"] > 0
    calls = len(ai_fake.Fake.requests)
    # the same stand-alone question in a new chat: our own answer cache, no AI call
    _, _, (_, cached) = ask(shop, a, "what are your timings")
    assert cached["outcome"] == "cached" and cached["text"].startswith("We are open")
    assert len(ai_fake.Fake.requests) == calls

    # retraining makes old cached answers stale
    shop.post(f"{AI}/agents/{a}/train", headers=WEB)
    ask(shop, a, "What are your timings?")
    assert len(ai_fake.Fake.requests) == calls + 1


def test_memory_handover_unknown_and_whatsapp_style(shop):
    a = agent_with_knowledge(shop)
    settings(shop, a, memory_messages=2)
    thread, _, _ = ask(shop, a, "Hi, my name is Ravi")
    for q in ("What are your timings?", "Is delivery free?", "What are your timings?"):
        ask(shop, a, q, thread)
    # older messages were folded into a summary that still knows the name
    db = SessionLocal()
    try:
        summary = db.get(AiThread, thread).summary
    finally:
        db.close()
    assert summary and "Ravi" in summary
    _, _, (_, name) = ask(shop, a, "What is my name?", thread)
    assert name["text"] == "Your name is Ravi."

    _, streamed, (_, hand) = ask(shop, a, "Can I talk to a human please?")
    assert hand["outcome"] == "handover" and hand["handover"] is True
    assert "[HANDOVER]" not in streamed and "[HANDOVER]" not in hand["text"]
    _, _, (_, unknown) = ask(shop, a, "Do you sell xylophones?")
    assert unknown["outcome"] == "unknown" and "[UNKNOWN]" not in unknown["text"]
    misses = shop.get(f"{AI}/agents/{a}/unanswered").json()["data"]
    assert {m["question"] for m in misses} >= {"Do you sell xylophones?", "Can I talk to a human please?"}

    ask(shop, a, "What are your timings?", channel="whatsapp")
    assert "This is WhatsApp" in ai_fake.Fake.requests[-1]["body"]["system"][0]["text"]

    threads = shop.get(f"{AI}/agents/{a}/threads", params={"channel": "test"}).json()
    assert threads["meta"]["total"] >= 4
    detail = shop.get(f"{AI}/agents/{a}/threads/{thread}").json()["data"]
    reply = [m for m in detail["items"] if m["role"] == "assistant"][-1]
    assert shop.post(f"{AI}/agents/{a}/messages/{reply['id']}/feedback", headers=WEB,
                     json={"value": "up"}).json()["data"]["feedback"] == "up"
    usage = shop.get(f"{AI}/agents/{a}/usage").json()["data"]
    assert usage["replies"] >= 7 and usage["outcomes"]["handover"] >= 1 and usage["feedback"]["up"] == 1
    assert usage["tokens"]["input"] > 0 and usage["cost"] > 0 and len(usage["daily"]) == 30


def test_search_mode_sends_only_matching_passages(shop):
    a = agent_with_knowledge(shop, knowledge_mode="search")
    _, _, (_, done) = ask(shop, a, "Is delivery free in Dehradun?")
    assert done["mode"] == "search" and "free above Rs 500" in done["text"]
    body = ai_fake.Fake.requests[-1]["body"]
    assert "Delivery is free" not in body["system"][0]["text"]  # not in the fixed part
    assert "<passage source=\"Delivery\">" in body["system"][1]["text"]


def test_openai_compatible_provider_and_errors(shop):
    conn = connect(shop, "gemini", "AIza-test-0000000000")
    a = shop.post(f"{AI}/agents", headers=WEB, json={"name": "Gem", "connection_id": conn["id"],
                                                     "model": "gpt-test-mini"}).json()["data"]["id"]
    shop.post(f"{AI}/agents/{a}/sources/faq", headers=WEB, json={"question": "Do you have parking?",
                                                                 "answer": "Yes, free parking behind the shop."})
    shop.post(f"{AI}/agents/{a}/train", headers=WEB)
    _, streamed, (_, done) = ask(shop, a, "Is there parking?")
    assert done["text"] == "Yes, free parking behind the shop." == streamed
    sent = ai_fake.Fake.requests[-1]
    assert sent["path"] == "/v1/chat/completions" and sent["body"]["messages"][0]["role"] == "system"
    assert sent["body"]["stream"] is True and sent["headers"]["Authorization"] == "Bearer AIza-test-0000000000"

    ai_fake.Fake.fail = "quota"
    _, _, (kind, err) = ask(shop, a, "Is there parking today?")
    assert kind == "error" and err["kind"] == "quota"
    ai_fake.Fake.fail = "key"
    _, _, (kind, err) = ask(shop, a, "Hello?")
    assert kind == "error" and err["kind"] == "key"
    assert shop.get(f"{AI}/connections").json()["data"][0]["status"] == "failed"


def test_monthly_budget_stops_the_agent(shop):
    a = agent_with_knowledge(shop)
    ask(shop, a, "Is delivery free?")
    settings(shop, a, monthly_budget=0.000001)
    _, _, (kind, err) = ask(shop, a, "What are your timings today?")
    assert kind == "error" and "budget" in err["message"]


# WhatsApp ---------------------------------------------------------------------------------------------

def test_whatsapp_answers_pauses_for_team_and_hands_over(shop):
    a = agent_with_knowledge(shop)
    shop.post(f"{BASE}/automation/rules", headers=WEB, json={"name": "Hi", "trigger": "welcome",
                                                              "reply": {"type": "text", "text": "Welcome!"}})
    shop.post(f"{BASE}/automation/rules", headers=WEB, json={"name": "Menu", "keywords": ["menu"],
                                                              "reply": {"type": "text", "text": "Our menu: ..."}})
    channels(shop, a, whatsapp_enabled=True)
    wa = "919811122233"
    inbound(shop, wa_id=wa, text="What are your timings?", name="Neha")
    assert texts_to(wa) == ["We are open from 9 am to 9 pm, all seven days."]  # no welcome: the AI greets
    conv = first_conversation(shop)
    msgs = shop.get(f"{BASE}/inbox/conversations/{conv['id']}/messages").json()["data"]
    assert msgs[-1]["sent_by"] == "AI agent" and msgs[-1]["auto_reply"] == "AI agent"
    state = shop.get(f"{AI}/chat/{conv['id']}").json()["data"]
    assert state["state"] == "active" and state["agent"]["name"] == "Asha"

    inbound(shop, wa_id=wa, text="menu")  # keyword replies still come first
    assert texts_to(wa)[-1] == "Our menu: ..."

    # a person from the team replies: the agent stays quiet
    shop.post(f"{BASE}/inbox/conversations/{conv['id']}/messages", headers=WEB, json={"text": "Hi Neha, Sumit here"})
    before = len(texts_to(wa))
    inbound(shop, wa_id=wa, text="Is delivery free?")
    assert len(texts_to(wa)) == before
    assert shop.get(f"{AI}/chat/{conv['id']}").json()["data"]["state"] == "team"
    # turned back on by hand
    assert shop.post(f"{AI}/chat/{conv['id']}/resume", headers=WEB).status_code == 200
    inbound(shop, wa_id=wa, text="Is delivery free?")
    assert "free above Rs 500" in texts_to(wa)[-1]

    # asks for a person: handed over, assigned, and the agent stays quiet until the chat is resolved
    inbound(shop, wa_id=wa, text="I want to talk to a human")
    assert texts_to(wa)[-1] == "Sure, connecting you with our team now."
    assert shop.get(f"{AI}/chat/{conv['id']}").json()["data"]["state"] == "handover"
    count = len(texts_to(wa))
    inbound(shop, wa_id=wa, text="hello?")
    assert len(texts_to(wa)) == count
    shop.put(f"{BASE}/inbox/conversations/{conv['id']}/status", headers=WEB, json={"status": "resolved"})
    inbound(shop, wa_id=wa, text="What are your timings?")
    assert texts_to(wa)[-1].startswith("We are open")

    # paused for one chat
    shop.post(f"{AI}/chat/{conv['id']}/pause", headers=WEB)
    count = len(texts_to(wa))
    inbound(shop, wa_id=wa, text="Is delivery free?")
    assert len(texts_to(wa)) == count
    convo = shop.get(f"{AI}/agents/{a}/threads", params={"channel": "whatsapp"}).json()["data"]
    assert convo and convo[0]["conversation_id"] == conv["id"]


def test_whatsapp_photo_gets_media_reply_and_off_hours_mode(shop):
    a = agent_with_knowledge(shop)
    channels(shop, a, whatsapp_enabled=True)
    wa = "919811100099"
    inbound(shop, wa_id=wa, type="image", image={"id": "IMG1", "mime_type": "image/jpeg"})
    assert texts_to(wa) == ["I can read text messages only. Please type your question."]
    # outside-hours mode: during office hours (always open here) the agent leaves the chat to the team
    shop.put(f"{BASE}/automation/settings", headers=WEB, json={"hours": {d: {"open": True, "start": "00:00",
                                                                              "end": "23:59"}
                                                                          for d in ("mon", "tue", "wed", "thu", "fri",
                                                                                    "sat", "sun")}})
    channels(shop, a, whatsapp={"mode": "outside_hours"})
    inbound(shop, wa_id="919811100098", text="What are your timings?")
    assert texts_to("919811100098") == []


# Website widget and API -------------------------------------------------------------------------------

def test_website_widget_lead_history_and_domains(shop):
    a = agent_with_knowledge(shop)
    key = shop.get(f"{AI}/agents/{a}").json()["data"]["public_key"]
    assert shop.get(f"{PUB}/{key}/config").status_code == 404  # widget off
    channels(shop, a, widget_enabled=True, widget={"allowed_domains": ["sumitsweets.com", "*.sumit.in"],
                                                   "lead": "required", "color": "#ff5500"})
    assert shop.get(f"{PUB}/{key}/config", params={"host": "evil.com"}).status_code == 403
    cfg = shop.get(f"{PUB}/{key}/config", params={"host": "shop.sumit.in"}).json()["data"]
    assert cfg["color"] == "#ff5500" and cfg["online"] is True and "allowed_domains" not in cfg
    assert cfg["whatsapp_url"] == "https://wa.me/15550000000"

    host = "sumitsweets.com"
    blocked = shop.post(f"{PUB}/{key}/chat", json={"message": "Hi", "host": host})
    assert blocked.status_code == 422 and blocked.json()["errors"]["lead"]
    bad_phone = shop.post(f"{PUB}/{key}/lead", json={"name": "Kiran", "phone": "12", "host": host})
    assert bad_phone.status_code == 422
    token = shop.post(f"{PUB}/{key}/lead", json={"name": "Kiran", "phone": "9876500011", "host": host}).json()
    token = token["data"]["token"]
    contact = shop.get(f"{BASE}/contacts", params={"search": "9876500011"}).json()["data"][0]
    assert contact["name"] == "Kiran" and "Website chat" in contact["tags"]

    events = sse(shop.post(f"{PUB}/{key}/chat", json={"message": "Is delivery free?", "token": token, "host": host,
                                                      "page": "https://sumitsweets.com/menu"}))
    assert events[0] == ("thread", {"token": token})
    done = events[-1][1]
    assert "free above Rs 500" in done["text"]
    history = shop.get(f"{PUB}/{key}/thread", params={"token": token, "host": host}).json()["data"]
    assert [m["role"] for m in history["items"]] == ["user", "assistant"] and history["visitor"]["name"] == "Kiran"
    assert shop.post(f"{PUB}/{key}/feedback", json={"token": token, "message_id": done["id"],
                                                     "value": "down"}).status_code == 200
    other = shop.post(f"{PUB}/{key}/feedback", json={"token": "nope", "message_id": done["id"], "value": "up"})
    assert other.status_code == 404
    web = shop.get(f"{AI}/agents/{a}/threads", params={"channel": "web"}).json()["data"]
    assert web[0]["visitor"]["name"] == "Kiran" and web[0]["visitor"]["page"].endswith("/menu")


def test_server_api_needs_the_secret_key(shop):
    a = agent_with_knowledge(shop)
    key = shop.get(f"{AI}/agents/{a}").json()["data"]["public_key"]
    turned_on = shop.put(f"{AI}/agents/{a}/channels", headers=WEB, json={"api_enabled": True})
    assert turned_on.status_code == 422  # needs a key first
    secret = shop.post(f"{AI}/agents/{a}/secret", headers=WEB).json()["data"]["secret"]
    assert secret.startswith("sk_ai_") and secret not in json.dumps(shop.get(f"{AI}/agents/{a}").json())
    channels(shop, a, api_enabled=True)
    body = {"message": "My name is Meera", "user": {"phone": "+919800000001"}}
    assert shop.post(f"{PUB}/{key}/messages", json=body).status_code == 401
    assert shop.post(f"{PUB}/{key}/messages", json=body, headers={"Authorization": "Bearer wrong"}).status_code == 401
    first = shop.post(f"{PUB}/{key}/messages", json=body, headers={"Authorization": f"Bearer {secret}"}).json()["data"]
    assert first["thread"] and first["outcome"] == "answered"
    again = shop.post(f"{PUB}/{key}/messages", json={"message": "What is my name?", "thread": first["thread"]},
                      headers={"Authorization": f"Bearer {secret}"}).json()["data"]
    assert again["reply"] == "Your name is Meera." and again["thread"] == first["thread"]


def test_duplicate_rekey_and_delete(shop):
    a = agent_with_knowledge(shop)
    dup = shop.post(f"{AI}/agents/{a}/duplicate", headers=WEB).json()["data"]
    assert dup["name"] == "Asha (copy)" and dup["knowledge"]["version"] == 1 and dup["knowledge"]["tokens"] > 0
    new_key = shop.post(f"{AI}/agents/{a}/public-key", headers=WEB).json()["data"]["public_key"]
    assert new_key.startswith("pk_")
    assert shop.delete(f"{AI}/agents/{dup['id']}", headers=WEB).status_code == 200
    assert len(shop.get(f"{AI}/agents").json()["data"]) == 1


def test_custom_addresses_must_be_public_in_production(shop, monkeypatch):
    from app.core.config import settings

    monkeypatch.setattr(type(settings), "is_production", property(lambda self: True))
    res = shop.post(f"{AI}/connections", headers=WEB, json={"provider": "custom", "api_key": "x",
                                                            "base_url": "http://127.0.0.1:11434/v1"})
    assert res.status_code == 422 and "public" in res.json()["errors"]["base_url"][0]
