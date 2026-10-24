"""Spreadsheets of contacts: read Excel or CSV, match columns, and check every row before anything is saved.

The same check powers the import preview (nothing written) and the import itself, so what people see in the
preview is exactly what happens.
"""
import csv
import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime

from fastapi import UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.modules.customer.account.models import Customer
from app.modules.customer.auth.service import field_error
from app.modules.customer.contacts.schemas import clean_attributes, clean_tags
from app.modules.customer.inbox.models import WaContact
from app.services.sms import config as sms_config
from app.services.sms import normalize_phone

MAX_ROWS = 10_000
MAX_BYTES = 5 * 1024 * 1024
FIELDS = ("phone", "name", "email", "tags")
GUESSES = {
    "phone": ("phone", "mobile", "number", "whatsapp", "wa_id", "phone_number", "mobile_number", "mobile_no",
              "phone_no", "contact", "contact_number", "whatsapp_number", "cell", "msisdn"),
    "name": ("name", "full_name", "contact_name", "customer_name", "customer", "client_name", "client"),
    "email": ("email", "email_address", "e-mail", "mail", "email_id"),
    "tags": ("tags", "tag", "labels", "label", "groups", "group"),
}
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")
SCIENTIFIC_RE = re.compile(r"^\s*\+?\d(\.\d+)?[eE]\+?\d+\s*$")


@dataclass
class Sheet:
    filename: str
    headers: list[str]
    rows: list[tuple[int, list[str]]]  # (row number as the person sees it in Excel, cell texts)
    sheet_name: str | None = None
    truncated: bool = False


@dataclass
class Mapping:
    phone: str | None
    name: str | None = None
    email: str | None = None
    tags: str | None = None
    skip: list[str] = field(default_factory=list)

    def custom(self, headers: list[str]) -> list[str]:
        used = {self.phone, self.name, self.email, self.tags, *self.skip}
        return [h for h in headers if h not in used]

    def as_dict(self) -> dict:
        return {"phone": self.phone, "name": self.name, "email": self.email, "tags": self.tags, "skip": self.skip}


# reading ----------------------------------------------------------------------------------

def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 919876543210.0 -> 919876543210
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M") if (value.hour, value.minute) != (0, 0) else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.startswith('="') and text.endswith('"'):  # our own export keeps Excel from mangling numbers
        text = text[2:-1]
    return text


def _headers(raw: list[str]) -> list[str]:
    out: list[str] = []
    for i, h in enumerate(raw):
        name = h.strip() or f"Column {_letter(i)}"
        base, n = name, 2
        while name in out:
            name, n = f"{base} ({n})", n + 1
        out.append(name)
    return out


def _letter(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _table(filename: str, grid, sheet_name: str | None = None) -> Sheet:
    header: list[str] | None = None
    rows: list[tuple[int, list[str]]] = []
    truncated = False
    for number, raw in enumerate(grid, start=1):
        cells = [_cell(v) for v in raw]
        if not any(cells):
            continue
        if header is None:
            header = _headers(cells)
            continue
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        cells = (cells + [""] * len(header))[: len(header)]
        rows.append((number, cells))
    if header is None:
        raise field_error(422, "file", "This file is empty. Put column names in the first row and one person per row.")
    return Sheet(filename=filename, headers=header, rows=rows, sheet_name=sheet_name, truncated=truncated)


def _read_xlsx(filename: str, raw: bytes) -> Sheet:
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - any parser failure means the file isn't a usable workbook
        raise field_error(422, "file", "We couldn't open this Excel file. Save it again as .xlsx or .csv and retry.")
    try:
        # the first sheet that has something in it
        for ws in book.worksheets:
            grid = list(ws.iter_rows(values_only=True))
            if any(any(c not in (None, "") for c in row) for row in grid):
                return _table(filename, grid, ws.title)
        raise field_error(422, "file", "This file is empty. Put column names in the first row and one person per row.")
    finally:
        book.close()


def _read_csv(filename: str, raw: bytes) -> Sheet:
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1252", errors="replace")  # files saved by older Excel on Windows
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return _table(filename, csv.reader(io.StringIO(text), dialect=dialect))


def read(file: UploadFile) -> Sheet:
    filename = file.filename or "contacts"
    raw = file.file.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise field_error(422, "file", "The file is larger than 5 MB. Split it into smaller files.")
    lower = filename.lower()
    if lower.endswith(".xls"):
        raise field_error(422, "file", "Old .xls files aren't supported. In Excel use Save as → Excel Workbook (.xlsx).")
    if lower.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK":  # PK: a zip, i.e. a renamed workbook
        return _read_xlsx(filename, raw)
    if lower.endswith((".csv", ".txt", ".tsv")):
        return _read_csv(filename, raw)
    raise field_error(422, "file", "Choose an Excel (.xlsx) or CSV file.")


# columns ----------------------------------------------------------------------------------

def _norm(header: str) -> str:
    return re.sub(r"[^a-z0-9@-]+", "_", header.strip().lower()).strip("_")


def guess(sheet: Sheet) -> dict:
    found: dict[str, str | None] = {}
    taken: set[str] = set()
    for key in FIELDS:
        found[key] = next((h for h in sheet.headers if h not in taken and _norm(h) in GUESSES[key]), None)
        if found[key]:
            taken.add(found[key])
    if not found["phone"]:
        # no telling header: take the column where most cells look like phone numbers
        best, score = None, 0.0
        sample = sheet.rows[:200]
        for i, h in enumerate(sheet.headers):
            if h in taken or not sample:
                continue
            hits = sum(1 for _, cells in sample if re.fullmatch(r"[\s+()\-.\d]{8,20}", cells[i] or "") and
                       len(re.sub(r"\D", "", cells[i])) >= 8)
            if hits / len(sample) > score:
                best, score = h, hits / len(sample)
        if score >= 0.6:
            found["phone"] = best
    return found


def mapping_from(sheet: Sheet, raw: str | None, require_phone: bool = True) -> Mapping:
    """The columns the person chose (JSON from the form), or our guess when they haven't chosen yet.

    The preview passes require_phone=False so a sheet whose phone column we can't spot still shows up,
    and the person can pick the column."""
    chosen: dict = {}
    if raw:
        try:
            chosen = json.loads(raw)
        except ValueError:
            raise field_error(422, "mapping", "Column choices are not valid")
        if not isinstance(chosen, dict):
            raise field_error(422, "mapping", "Column choices are not valid")
    guessed = guess(sheet)
    values = {k: (chosen[k] or None) if k in chosen else guessed[k] for k in FIELDS}
    for key, header in values.items():
        if header is not None and header not in sheet.headers:
            raise field_error(422, f"mapping.{key}", f"There is no column called “{header}” in this file")
    picked = [h for h in values.values() if h]
    if len(picked) != len(set(picked)):
        raise field_error(422, "mapping", "Use each column for one thing only")
    if not values["phone"] and require_phone:
        raise field_error(422, "mapping.phone", "Choose the column that has the phone numbers")
    skip = [h for h in chosen.get("skip") or [] if h in sheet.headers]
    return Mapping(phone=values["phone"], name=values["name"], email=values["email"], tags=values["tags"],
                   skip=skip)


# checking ---------------------------------------------------------------------------------

def _split_tags(text: str) -> list[str]:
    return [t for t in re.split(r"[;,|]", text or "") if t.strip()]


def check(db: Session, customer: Customer, sheet: Sheet, mapping: Mapping, update_existing: bool,
          tag: str | None = None) -> dict:
    """Every row with what will happen to it. Nothing is written."""
    idx = {h: i for i, h in enumerate(sheet.headers)}
    custom = mapping.custom(sheet.headers)
    extra_tag = clean_tags([tag])[0] if tag and clean_tags([tag]) else None
    country = sms_config.load(db)["default_country_code"]
    existing = {
        c.wa_id: c for c in db.query(WaContact).filter(WaContact.customer_id == customer.id).all()
    }
    first_row: dict[str, int] = {}
    rows: list[dict] = []
    counts = {"total": 0, "new": 0, "update": 0, "exists": 0, "duplicate": 0, "invalid": 0, "warnings": 0}

    def get(cells: list[str], header: str | None) -> str:
        return cells[idx[header]].strip() if header else ""

    for number, cells in sheet.rows:
        counts["total"] += 1
        raw_phone = get(cells, mapping.phone)
        name = get(cells, mapping.name)[:150]
        email = get(cells, mapping.email).lower()[:255]
        tags = clean_tags([*_split_tags(get(cells, mapping.tags)), *([extra_tag] if extra_tag else [])])
        fields = clean_attributes({h: cells[idx[h]] for h in custom})
        warnings: list[str] = []
        row = {"row": number, "phone": raw_phone, "wa_id": None, "name": name or None, "email": email or None,
               "tags": tags, "fields": fields, "status": "new", "message": None, "duplicate_of": None,
               "warnings": warnings, "contact_id": None}
        rows.append(row)

        if not mapping.phone:
            row.update(status="invalid", message="Choose the column that has the phone numbers")
        elif not raw_phone:
            row.update(status="invalid", message="No phone number")
        elif SCIENTIFIC_RE.match(raw_phone):
            row.update(status="invalid",
                       message=f"Excel shortened this number to {raw_phone}. Format the column as Text and re-type it")
        else:
            phone = normalize_phone(raw_phone, country)
            if not phone:
                row.update(status="invalid", message="Not a valid phone number")
            else:
                row["wa_id"] = phone[1:]
        if row["status"] == "invalid":
            counts["invalid"] += 1
            continue

        wa_id = row["wa_id"]
        if wa_id in first_row:
            row.update(status="duplicate", duplicate_of=first_row[wa_id],
                       message=f"Same number as row {first_row[wa_id]}")
            counts["duplicate"] += 1
            continue
        first_row[wa_id] = number

        if email and not EMAIL_RE.match(email):
            warnings.append(f"Email “{email}” isn't valid, so it's left out")
            row["email"] = None
        contact = existing.get(wa_id)
        if contact is not None:
            row["contact_id"] = contact.id
            row["status"] = "update" if update_existing else "exists"
            row["message"] = "Already a contact: details will be added" if update_existing else \
                "Already a contact: left as it is"
            if contact.blocked:
                warnings.append("Blocked: won't get messages")
            elif contact.opted_out:
                warnings.append("Opted out: won't get campaigns")
        counts[row["status"]] += 1
        if warnings:
            counts["warnings"] += 1

    return {
        "filename": sheet.filename,
        "sheet": sheet.sheet_name,
        "headers": sheet.headers,
        "mapping": mapping.as_dict(),
        "custom": [{"column": h, "field": next(iter(clean_attributes({h: "x"})), None)} for h in custom],
        "counts": counts,
        "truncated": sheet.truncated,
        "max_rows": MAX_ROWS,
        "rows": rows,
    }


def save(db: Session, customer: Customer, report: dict) -> dict:
    """Write what `check` decided: new contacts added, existing ones filled in (when asked)."""
    existing = {
        c.wa_id: c for c in db.query(WaContact).filter(WaContact.customer_id == customer.id).all()
    }
    created = updated = 0
    for row in report["rows"]:
        if row["status"] == "new":
            db.add(WaContact(customer_id=customer.id, wa_id=row["wa_id"], source="import", name=row["name"],
                             email=row["email"], tags=row["tags"], attributes=row["fields"]))
            created += 1
        elif row["status"] == "update":
            contact = existing[row["wa_id"]]
            contact.name = row["name"] or contact.name
            contact.email = row["email"] or contact.email
            contact.tags = clean_tags([*(contact.tags or []), *row["tags"]])
            contact.attributes = clean_attributes({**(contact.attributes or {}), **row["fields"]})
            updated += 1
    db.commit()
    return {"created": created, "updated": updated}


# sample -----------------------------------------------------------------------------------

SAMPLE_HEADERS = ["Phone", "Name", "Email", "Tags", "City", "Order ID", "Amount"]
SAMPLE_ROWS = [
    ["+91 98765 43210", "Asha Verma", "asha@example.com", "VIP; Lead", "Delhi", "ORD-1001", "1499"],
    ["9876543211", "Ravi Kumar", "", "Wholesale", "Pune", "ORD-1002", "899"],
    ["+44 7700 900123", "Tom Baker", "tom@example.co.uk", "", "London", "ORD-1003", "2350"],
]


def sample(fmt: str) -> StreamingResponse:
    if fmt == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(SAMPLE_HEADERS)
        writer.writerows(SAMPLE_ROWS)
        return StreamingResponse(iter(["﻿" + buf.getvalue()]), media_type="text/csv; charset=utf-8",
                                 headers={"Content-Disposition": 'attachment; filename="contacts-sample.csv"'})

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    ws = book.active
    ws.title = "Contacts"
    ws.append(SAMPLE_HEADERS)
    for row in SAMPLE_ROWS:
        ws.append(row)
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F766E")
    for cell in ws[1]:
        cell.font, cell.fill = head, fill
        cell.alignment = Alignment(vertical="center")
    # phone numbers as text, so Excel never turns them into 9.18E+11
    for row in ws.iter_rows(min_row=1, max_row=1000, min_col=1, max_col=1):
        for cell in row:
            cell.number_format = "@"
    for letter, width in zip("ABCDEFG", (20, 18, 24, 16, 12, 12, 10)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    notes = book.create_sheet("How to fill")
    for line in [
        ["How to fill this sheet"],
        [],
        ["Phone", "Required. With country code (+91 98765 43210) or 10 digits for your default country."],
        ["Name", "Optional. Used for {{name}} in messages."],
        ["Email", "Optional."],
        ["Tags", "Optional. Separate several with ; e.g. VIP; Lead"],
        ["Any other column", "Saved on the contact as a custom field (e.g. City, Order ID) and can be used in "
                             "campaign messages."],
        [],
        ["Tip", "Keep the Phone column formatted as Text so Excel doesn't shorten long numbers."],
    ]:
        notes.append(line)
    notes["A1"].font = Font(bold=True, size=13)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 90
    out = io.BytesIO()
    book.save(out)
    return StreamingResponse(
        iter([out.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="contacts-sample.xlsx"'},
    )
